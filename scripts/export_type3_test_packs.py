#!/usr/bin/env python3
"""Export Type 3 Test 1 packs — AI-1 and AI-2 only (Field / Value / Note)."""

from __future__ import annotations

import json
import re
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from ai_agent.api.lab import load_lab_config
from ai_agent.v2.hearing_parser import parse_hearing_sheet
from ai_agent.v2.page_catalog import nested_fields_for_section
from ai_agent.v2.prompt_packs import (
    default_ai1_planner_for_type,
    default_ai2_system_for_type,
)

ROOT = Path(__file__).resolve().parents[1]
HEARING = ROOT / "demo/v2/samples/type3-satellite.csv"
# Prefer latest regenerated content export
CONTENT_CANDIDATES = [
    ROOT / "有限会社太陽塗装-sections-2026-09-17 (1).xlsx",
    ROOT / "有限会社太陽塗装-sections-2026-09-17.xlsx",
]
OUT_AI1 = ROOT / "Type3-Test1-AI1-Sections.xlsx"
OUT_AI2 = ROOT / "Type3-Test1-AI2-Contents.xlsx"
OUT_AI1_JSON = ROOT / "Type3-Test1-AI1-structure.json"

HEARING_URL = (
    "https://drive.google.com/file/d/1NmQcmogLrCzpBpFMWXC_NpEqg2wA7jNp/view?usp=sharing"
)
CLIENT = "有限会社太陽塗装"
HEARING_FILE = "type3-satellite.csv"
AI1_MODEL = "gemini-3.5-flash-paid"
AI2_MODEL = "gemini-3.5-flash-lite-paid"

KEY_MAP = {
    "home": "TOP",
    "concept": "コンセプト",
    "service": "サービス",
    "faq": "よくある質問",
    "greeting": "代表あいさつ",
    "access": "アクセス",
    "blog": "ブログ",
    "contact": "お問い合わせ",
    "sitemap": "サイトマップ",
    "privacy": "プライバシーポリシー",
    "reviews": "お客様の声",
    "ai-blog": "AIブログ",
    "column": "コラム",
}


def _parse_body(raw: object) -> object:
    if raw is None:
        return ""
    t = str(raw).strip()
    if not t:
        return ""
    if t.startswith("{") or t.startswith("["):
        try:
            return json.loads(t)
        except Exception:
            return t
    return t


def load_contents_json(path: Path) -> OrderedDict[str, OrderedDict[str, object]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    pages_ws = wb["Pages"]
    order: list[tuple[str, str, str]] = []  # key, page_id, group
    for r in range(3, pages_ws.max_row + 1):
        group = str(pages_ws.cell(r, 1).value or "")
        page_id = str(pages_ws.cell(r, 2).value or "")
        slug = str(pages_ws.cell(r, 3).value or "")
        if not page_id:
            continue
        if group == "SEO":
            key = f"SEO:{slug or page_id}"
        elif group == "タグ":
            key = f"TAG:{slug or page_id}"
        else:
            key = KEY_MAP.get(slug or page_id, slug or page_id)
        order.append((key, page_id, group))

    all_ws = wb["All"]
    headers = [all_ws.cell(2, c).value for c in range(1, all_ws.max_column + 1)]
    by_page: dict[str, OrderedDict[str, object]] = {}
    for r in range(3, all_ws.max_row + 1):
        row = {headers[c - 1]: all_ws.cell(r, c).value for c in range(1, all_ws.max_column + 1)}
        pid = str(row.get("ページID") or "")
        sid = str(row.get("セクションID") or "")
        if not pid or not sid:
            continue
        by_page.setdefault(pid, OrderedDict())[sid] = _parse_body(row.get("本文"))

    out: OrderedDict[str, OrderedDict[str, object]] = OrderedDict()
    for key, page_id, _group in order:
        out[key] = by_page.get(page_id, OrderedDict())
    return out


def build_ai1_structure_from_contents(
    contents: OrderedDict[str, OrderedDict[str, object]],
) -> OrderedDict[str, list[dict[str, str]]]:
    """AI-1 result shape: per page list of {id,label,mode,rule} (no Japanese copy)."""
    final: OrderedDict[str, list[dict[str, str]]] = OrderedDict()
    for page_key, sections in contents.items():
        rows: list[dict[str, str]] = []
        for sid in sections.keys():
            fields = nested_fields_for_section({"id": sid})
            if fields:
                shape = "{" + ",".join(fields) + "}"
                rule = f"Fill nested {shape} from hearing facts only; empty if missing."
                mode = "facts" if sid in {"business_info", "cta", "access_details", "contact_details", "keyword"} else "generate"
            else:
                rule = "Fill from hearing facts only; empty string if missing."
                mode = "facts" if sid in {"keyword", "reference_url"} else "generate"
            if sid in {"faq_items", "review_items"} or sid.startswith("review_"):
                mode = "blank"
                rule = "Blank — hearing has no Q&A / review bodies."
            if page_key in {"代表あいさつ", "お客様の声"}:
                mode = "blank"
                rule = "Blank — hearing has no staff / review bodies."
            if page_key in {"ブログ", "AIブログ", "コラム", "サイトマップ", "プライバシーポリシー"} and sid == "listing_intro":
                mode = "shell"
                rule = "Listing shell only from hearing site name; no invented articles."
            rows.append(
                {
                    "id": sid,
                    "label": sid.replace("_", " ").title(),
                    "mode": mode,
                    "rule": rule,
                }
            )
        final[page_key] = rows
    return final


def write_pack(
    path: Path,
    *,
    title: str,
    test_name: str,
    model: str,
    model_note: str,
    system_prompt: str,
    result_json: str,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test1"
    ws.append([title])
    ws.append(["Field", "Value", "Note"])
    rows = [
        ("Test name", test_name, ""),
        ("Site / client", CLIENT, ""),
        ("AI model", model, model_note),
        ("Input data (hearing file name)", HEARING_FILE, ""),
        ("Input data (hearing URL)", HEARING_URL, ""),
        ("System prompt (FULL)", system_prompt, ""),
        ("Result (FULL JSON)", result_json, ""),
    ]
    for a, b, c in rows:
        ws.append([a, b, c])
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 28
    # wrap prompt + result
    for row_idx in (8, 9):  # System prompt, Result
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 220
    wb.save(path)


def main() -> None:
    content_path = next((p for p in CONTENT_CANDIDATES if p.exists()), None)
    if not content_path:
        raise SystemExit("No content export xlsx found")

    cfg = load_lab_config()
    tp = (cfg.get("type_prompts") or {}).get("type3") or {}
    planner = str(tp.get("planner_system_prompt") or "").strip() or default_ai1_planner_for_type("type3")
    writer = str(tp.get("system_prompt") or "").strip() or default_ai2_system_for_type("type3")

    # Ensure hearing parses (category attach etc.)
    parse_hearing_sheet(HEARING.read_text(encoding="utf-8-sig"))

    contents = load_contents_json(content_path)
    # Strip accidental _note if present
    for page in contents.values():
        page.pop("_note", None)

    ai1_struct = build_ai1_structure_from_contents(contents)
    ai1_json = json.dumps(ai1_struct, ensure_ascii=False, indent=4)
    ai2_json = json.dumps(contents, ensure_ascii=False, indent=4)

    if re.search(r'"hero"|\"hero_', ai1_json) or re.search(r'"hero"|\"hero_', ai2_json):
        raise SystemExit("refusing: hero keys found")

    write_pack(
        OUT_AI1,
        title="Test 1 — Type 3 — AI-1 Sections",
        test_name="Test 1 — Type 3 Satellite — AI-1 Sections",
        model=AI1_MODEL,
        model_note="AI-1 = Sections Planner",
        system_prompt=planner,
        result_json=ai1_json,
    )
    write_pack(
        OUT_AI2,
        title="Test 1 — Type 3 — AI-2 Contents",
        test_name="Test 1 — Type 3 Satellite — AI-2 Contents",
        model=AI2_MODEL,
        model_note="AI-2 = Content Writer",
        system_prompt=writer,
        result_json=ai2_json,
    )
    OUT_AI1_JSON.write_text(ai1_json + "\n", encoding="utf-8")

    # Remove old combined review pack if present
    old = ROOT / "Type3-Test1-client-review.xlsx"
    if old.exists():
        old.unlink()

    print(f"wrote {OUT_AI1.name}")
    print(f"wrote {OUT_AI2.name}")
    print(f"wrote {OUT_AI1_JSON.name}")
    print(f"source contents: {content_path.name}")
    print(f"removed: {old.name}" if not old.exists() else "")


if __name__ == "__main__":
    main()
