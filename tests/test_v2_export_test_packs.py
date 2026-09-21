"""Live export uses hearing/category site+pages JSON (no id/label/mode/rule)."""

from __future__ import annotations

import json
from io import BytesIO

from ai_agent.v2.export import (
    build_ai1_result,
    build_ai2_result,
    build_test_packs_payload,
    page_export_key,
    packs_to_xlsx_bytes,
    sections_to_xlsx_bytes,
)


def test_page_export_key_nav_and_seo():
    assert page_export_key(group="page", slug="home", nav_label="TOP") == "TOP"
    assert page_export_key(group="page", slug="concept", nav_label="コンセプト") == "コンセプト"
    assert page_export_key(group="seo", slug="seo-1", nav_label="リフォーム") == "SEO:リフォーム"
    assert page_export_key(group="tag", slug="tag-1", nav_label="遮熱") == "TAG:遮熱"


def test_ai1_ai2_result_site_pages_no_meta():
    hearing = {
        "project": {
            "business_name": "有限会社太陽塗装",
            "purpose": "集客 (訪問型)",
            "production_kind": "通常",
            "domain": "taiyotoso.jp",
            "industry": "外壁塗装",
            "area": "長崎",
            "site_category": "lead_gen",
            "site_brief": {
                "category": "lead_gen",
                "purpose": "集客 (訪問型)",
                "production_kind": "通常",
                "domain": "taiyotoso.jp",
                "industry": "外壁塗装",
                "area": "長崎",
                "business_name": "有限会社太陽塗装",
                "audience": "customers",
                "goal": "service_inquiries",
            },
        },
        "site_category": "lead_gen",
        "store": {"phone": "0120-011-923"},
    }
    blueprint = {
        "site_name": "有限会社太陽塗装",
        "production_type": "type3",
        "pages": [
            {
                "id": "home",
                "slug": "home",
                "nav_label": "TOP",
                "type": "top_satellite",
                "sections": [
                    {
                        "id": "top_catchphrase",
                        "label": "Catch",
                        "mode": "generate",
                        "rule": "x",
                        "fields": ["brand_name", "catchphrase", "short_description"],
                    },
                    {
                        "id": "lead",
                        "label": "Lead",
                        "mode": "generate",
                        "rule": "x",
                        "fields": ["heading", "body"],
                    },
                ],
            },
            {
                "id": "concept",
                "slug": "concept",
                "nav_label": "コンセプト",
                "type": "concept",
                "sections": [
                    {
                        "id": "concept_catchphrase",
                        "label": "Catch",
                        "mode": "generate",
                        "rule": "x",
                        "fields": ["catchphrase", "short_description"],
                    },
                    {
                        "id": "point_1",
                        "label": "Point 1",
                        "mode": "expand",
                        "rule": "x",
                        "fields": ["title", "description"],
                    },
                ],
            },
        ],
        "seo_pages": [
            {
                "id": "seo-1",
                "slug": "seo-1",
                "nav_label": "リフォーム",
                "sections": [
                    {"id": "keyword", "mode": "facts", "rule": "x"},
                    {
                        "id": "seo_intro",
                        "mode": "generate",
                        "rule": "x",
                        "fields": ["heading", "body"],
                    },
                ],
            }
        ],
        "tag_pages": [],
    }
    rows = [
        {
            "group": "page",
            "page_slug": "home",
            "nav_label": "TOP",
            "section_id": "top_catchphrase",
            "text": json.dumps(
                {"brand_name": "有限会社太陽塗装", "catchphrase": "A", "short_description": "B"},
                ensure_ascii=False,
            ),
        },
        {
            "group": "page",
            "page_slug": "home",
            "nav_label": "TOP",
            "section_id": "lead",
            "text": json.dumps({"heading": "H", "body": "リード文"}, ensure_ascii=False),
        },
        {
            "group": "page",
            "page_slug": "concept",
            "nav_label": "コンセプト",
            "section_id": "concept_catchphrase",
            "text": '{"catchphrase":"C","short_description":"D"}',
        },
        {
            "group": "seo",
            "page_slug": "seo-1",
            "nav_label": "リフォーム",
            "section_id": "keyword",
            "text": "リフォーム",
        },
    ]
    ai1 = build_ai1_result(blueprint, rows, hearing=hearing)
    assert set(ai1.keys()) == {"site", "pages"}
    assert ai1["site"]["category"] == "lead_gen"
    assert ai1["site"]["name"] == "有限会社太陽塗装"
    assert "TOP" in ai1["pages"]
    assert "id" not in ai1["pages"]["TOP"]
    top = ai1["pages"]["TOP"]
    assert isinstance(top, dict)
    assert "top_catchphrase" in top
    assert top["top_catchphrase"]["brand_name"] == "会社名"
    assert top["top_catchphrase"]["catchphrase"] == "短いキャッチ（15–28字）"
    assert "説明" in top["top_catchphrase"]["short_description"]
    assert "1文" in ai1["pages"]["コンセプト"]["point_1"]["description"]
    dumped = json.dumps(ai1, ensure_ascii=False)
    assert '"id"' not in dumped
    assert '"label"' not in dumped
    assert '"mode"' not in dumped
    assert '"rule"' not in dumped
    assert ai1["pages"]["SEO:リフォーム"]["keyword"] == "リフォーム"

    ai2 = build_ai2_result(rows, blueprint, hearing=hearing)
    assert ai2["site"]["category"] == "lead_gen"
    assert ai2["pages"]["TOP"]["top_catchphrase"]["catchphrase"] == "A"
    assert ai2["pages"]["TOP"]["lead"]["body"] == "リード文"
    assert ai2["pages"]["コンセプト"]["concept_catchphrase"]["catchphrase"] == "C"
    assert ai2["pages"]["SEO:リフォーム"]["keyword"] == "リフォーム"
    # AI-2 must not keep AI-1 planner hints on filled nested fields
    assert ai2["pages"]["TOP"]["top_catchphrase"]["brand_name"] == "有限会社太陽塗装"


def test_xlsx_pack_has_two_sheets():
    hearing = {
        "site_category": "lead_gen",
        "project": {
            "business_name": "Test Co",
            "purpose": "集客",
            "site_brief": {"category": "lead_gen", "purpose": "集客", "business_name": "Test Co"},
        },
    }
    blueprint = {
        "site_name": "Test Co",
        "production_label": "サテライト",
        "pages": [
            {
                "id": "home",
                "slug": "home",
                "nav_label": "TOP",
                "sections": [
                    {
                        "id": "lead",
                        "label": "Lead",
                        "mode": "generate",
                        "rule": "x",
                        "fields": ["heading", "body"],
                    }
                ],
            }
        ],
    }
    rows = [
        {
            "group": "page",
            "page_slug": "home",
            "nav_label": "TOP",
            "section_id": "lead",
            "text": json.dumps({"heading": "H", "body": "本文"}, ensure_ascii=False),
            "site": "Test Co",
        }
    ]
    packs = build_test_packs_payload(
        blueprint=blueprint,
        rows=rows,
        planner_prompt="PLANNER PROMPT",
        writer_prompt="WRITER PROMPT",
        ai1_model="model-a",
        ai2_model="model-b",
        hearing_file="type3-satellite.csv",
        hearing=hearing,
    )
    assert packs["ai1"]["values"][0][0].startswith("Live —")
    assert packs["ai1"]["values"][1] == ["Field", "Value", "Note"]
    assert packs["ai1"]["values"][8][0] == "Result"
    outline = packs["ai1"]["values"][8][1]
    assert "site" in outline and "pages" in outline and "TOP" in outline
    nbsp4 = "\u00A0" * 4
    assert (nbsp4 + "TOP\n") in outline or ("\n" + nbsp4 + "TOP\n") in outline
    assert any(r[1] for r in packs["ai1"]["result_runs"]), "expected bold runs"
    parsed = json.loads(packs["ai1"]["result_json"])
    assert "site" in parsed and "pages" in parsed
    assert "id" not in json.dumps(parsed.get("pages", {}))
    raw = packs_to_xlsx_bytes(packs)
    assert raw[:2] == b"PK"

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == ["AI-1 Sections", "AI-2 Contents"]
    ws = wb["AI-1 Sections"]
    assert ws["A2"].value == "Field"
    assert ws["B8"].value == "PLANNER PROMPT"
    assert "lead_gen" in str(ws["B9"].value) or "TOP" in str(ws["B9"].value)

    raw2 = sections_to_xlsx_bytes(
        rows,
        blueprint=blueprint,
        planner_prompt="P",
        writer_prompt="W",
        hearing_file="h.csv",
        hearing=hearing,
    )
    assert raw2[:2] == b"PK"
