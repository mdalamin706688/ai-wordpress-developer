"""Export v2 blueprint/sections — CSV (flat) + Test-pack xlsx (AI-1 / AI-2)."""

from __future__ import annotations

import csv
import io
import json
import re
from collections import OrderedDict
from datetime import UTC, datetime
from typing import Any

# Slug → Japanese page key when nav_label is missing (matches Test1 packs).
_SLUG_PAGE_KEY = {
    "home": "TOP",
    "concept": "コンセプト",
    "service": "サービス",
    "faq": "よくある質問",
    "greeting": "代表あいさつ",
    "access": "アクセス",
    "blog": "ブログ",
    "contact": "お問い合わせ",
    "sitemap": "サイトマップ",
    "privacy": "プライバシーポリシー",
    "reviews": "お客様の声",
    "ai-blog": "AIブログ",
    "column": "コラム",
    "menu": "料金表",
    "recruit": "求人",
}


def blueprint_to_section_rows(blueprint: dict[str, Any]) -> list[dict[str, str]]:
    """Flatten blueprint pages → export rows (before AI-2 content)."""
    rows: list[dict[str, str]] = []
    site = str(blueprint.get("site_name") or "")

    def add_page(page: dict[str, Any], *, group: str = "page") -> None:
        pid = str(page.get("id") or "")
        slug = str(page.get("slug") or "")
        for sec in page.get("sections") or []:
            if not isinstance(sec, dict):
                continue
            rows.append(
                {
                    "site": site,
                    "group": group,
                    "page_id": pid,
                    "page_slug": slug,
                    "page_type": str(page.get("type") or ""),
                    "section_id": str(sec.get("id") or ""),
                    "section_label": str(sec.get("label") or ""),
                    "mode": str(sec.get("mode") or ""),
                    "text": "",
                    "seeds": " | ".join(page.get("content_seeds") or [])[:500],
                    "nav_label": str(page.get("nav_label") or ""),
                }
            )

    for page in blueprint.get("pages") or []:
        if isinstance(page, dict):
            add_page(page, group="page")
    for page in blueprint.get("seo_pages") or []:
        if isinstance(page, dict):
            add_page(page, group="seo")
    for page in blueprint.get("tag_pages") or []:
        if isinstance(page, dict):
            add_page(page, group="tag")
    return rows


def sections_to_csv(rows: list[dict[str, str]]) -> str:
    fields = [
        "site",
        "group",
        "page_id",
        "page_slug",
        "page_type",
        "section_id",
        "section_label",
        "mode",
        "text",
        "seeds",
    ]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, lineterminator="\n")
    w.writeheader()
    for row in rows:
        w.writerow({k: row.get(k, "") for k in fields})
    return buf.getvalue()


def _group_norm(group: str) -> str:
    g = (group or "").lower().strip()
    if g in {"page", "nav", "ナビ"}:
        return "page"
    if g in {"seo", "ＳＥＯ"}:
        return "seo"
    if g in {"tag", "タグ"}:
        return "tag"
    return g or "page"


def page_export_key(
    *,
    group: str = "page",
    slug: str = "",
    page_id: str = "",
    nav_label: str = "",
) -> str:
    """Stable Japanese / SEO / TAG key used in Result JSON."""
    g = _group_norm(group)
    slug = str(slug or "").strip()
    page_id = str(page_id or "").strip()
    nav = str(nav_label or "").strip()
    if g == "seo":
        return f"SEO:{nav or slug or page_id or 'page'}"
    if g == "tag":
        return f"TAG:{nav or slug or page_id or 'page'}"
    if slug == "home" or nav.upper() == "TOP":
        return "TOP"
    if nav:
        return nav
    return _SLUG_PAGE_KEY.get(slug, slug or page_id or "page")


def _parse_body(raw: object) -> object:
    if raw is None:
        return ""
    if isinstance(raw, (dict, list)):
        return raw
    t = str(raw).strip()
    if not t:
        return ""
    if t.startswith("{") or t.startswith("["):
        try:
            return json.loads(t)
        except Exception:
            return t
    return t


def _type_label(blueprint: dict[str, Any]) -> str:
    raw = str(
        blueprint.get("production_label")
        or blueprint.get("production_type")
        or blueprint.get("type")
        or ""
    )
    low = raw.lower()
    if "type4" in low or "サテライトリニューアル" in raw:
        return "Type 4"
    if "type3" in low or "サテライト" in raw or "satellite" in low:
        return "Type 3"
    if "type2" in low or "リニューアル" in raw or "renewal" in low:
        return "Type 2"
    if "type1" in low or "新規" in raw:
        return "Type 1"
    return raw or "Type ?"


def _page_index_from_blueprint(blueprint: dict[str, Any]) -> OrderedDict[str, dict[str, Any]]:
    """Ordered page_key → page dict (+ group)."""
    out: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for group, key in (("page", "pages"), ("seo", "seo_pages"), ("tag", "tag_pages")):
        for page in blueprint.get(key) or []:
            if not isinstance(page, dict):
                continue
            pk = page_export_key(
                group=group,
                slug=str(page.get("slug") or ""),
                page_id=str(page.get("id") or ""),
                nav_label=str(page.get("nav_label") or ""),
            )
            if pk not in out:
                out[pk] = {**page, "_export_group": group}
    return out


def build_site_export_block(
    hearing: dict[str, Any] | None,
    *,
    blueprint: dict[str, Any] | None = None,
    site_name: str = "",
) -> OrderedDict[str, Any]:
    """Compact site / category block for Result JSON (hearing-driven)."""
    hearing = hearing if isinstance(hearing, dict) else {}
    bp = blueprint or {}
    project = hearing.get("project") if isinstance(hearing.get("project"), dict) else {}
    brief = hearing.get("site_brief") if isinstance(hearing.get("site_brief"), dict) else {}
    if not brief and isinstance(project.get("site_brief"), dict):
        brief = project.get("site_brief") or {}
    store = hearing.get("store") if isinstance(hearing.get("store"), dict) else {}

    refs: list[str] = []
    for item in hearing.get("reference_sites") or []:
        if isinstance(item, dict):
            url = str(item.get("url") or "").strip()
            if url:
                refs.append(url)
        elif str(item or "").strip():
            refs.append(str(item).strip())

    line_url = ""
    wg = hearing.get("writing_guidance") if isinstance(hearing.get("writing_guidance"), dict) else {}
    remarks = str(wg.get("remarks") or "")
    if "line.me" in remarks:
        m = re.search(r"https://line\.me/\S+", remarks)
        if m:
            line_url = m.group(0).rstrip("。．)）\"'")

    name = str(
        site_name
        or brief.get("business_name")
        or project.get("business_name")
        or store.get("name")
        or bp.get("site_name")
        or ""
    )
    out: OrderedDict[str, Any] = OrderedDict()
    out["name"] = name
    out["domain"] = str(brief.get("domain") or project.get("domain") or "")
    out["category"] = str(
        brief.get("category") or hearing.get("site_category") or project.get("site_category") or ""
    )
    out["purpose"] = str(brief.get("purpose") or project.get("purpose") or "")
    out["production_kind"] = str(brief.get("production_kind") or project.get("production_kind") or "")
    out["audience"] = str(brief.get("audience") or "")
    out["goal"] = str(brief.get("goal") or "")
    out["industry"] = str(brief.get("industry") or project.get("industry") or "")
    out["industry_category"] = str(project.get("industry_category") or "")
    out["area"] = str(brief.get("area") or project.get("area") or "")
    out["model"] = str(project.get("model") or "")
    out["ai_support"] = str(project.get("ai_support") or (hearing.get("ai_blog") or {}).get("ai_support") or "")
    out["reference_sites"] = refs
    out["line_url"] = line_url
    return out


def _ai1_field_hint(field: str, *, section_id: str = "", page_key: str = "") -> str:
    """Short planner hint for AI-1 Result (not final copy). description = one sentence."""
    f = str(field or "").strip().lower()
    sid = str(section_id or "").strip().lower()
    page = str(page_key or "")

    if f == "description":
        if sid.startswith("service_teaser_") or sid.startswith("service_"):
            return "ヒアリングのサービス内容から、何をするかを1文で書く。"
        if sid.startswith("selling_point_") or sid.startswith("point_"):
            return "ヒアリングの強み・コンセプト要点を、お客さま向けに1文で書く。"
        if sid.startswith("seo_point_") or sid.startswith("tag_point_"):
            return "このキーワードページ専用の論点を、TOPのブランド由来話を使わず1文で書く。"
        if "review" in sid:
            return "ヒアリングに口コミ本文がある場合のみ1文。なければ空。"
        return "ヒアリング事実だけを使い、このセクション内容を1文で書く。"

    short = {
        "brand_name": "会社名",
        "catchphrase": "短いキャッチ（15–28字）",
        "short_description": "補足の短い説明",
        "heading": "見出し",
        "body": "本文（短く）",
        "lead": "リード（短く）",
        "title": "タイトル",
        "label": "CTAラベル",
        "phone": "電話",
        "url": "URL",
        "line_url": "LINE URL",
        "methods": "連絡手段",
        "name": "名称",
        "postal": "郵便番号",
        "address": "住所",
        "hours": "営業時間",
        "closed": "定休日",
        "payment": "支払い",
        "email": "メール",
        "instagram": "Instagram",
        "station": "最寄駅",
        "parking": "駐車場",
        "map_note": "地図メモ（場所のみ）",
        "map_url": "地図URL",
        "form_note": "フォーム注意",
        "role": "役職",
        "message": "あいさつ文",
        "career": "経歴",
        "cta_label": "CTA文言",
        "keyword": page[4:] if page.startswith(("SEO:", "TAG:")) else "キーワード",
    }
    if f in short:
        return short[f]
    return "ヒアリングから短く"


def _ai1_section_hint_value(sec: dict[str, Any], *, page_key: str = "") -> object:
    """AI-1 nested shells with readable short hints (not empty, not final copy)."""
    from ai_agent.v2.page_catalog import nested_fields_for_section

    sid = str(sec.get("id") or "").strip()
    fields = nested_fields_for_section(sec)
    if fields:
        out: dict[str, str] = {}
        for f in fields:
            fl = str(f)
            if fl == "keyword" and page_key.startswith(("SEO:", "TAG:")):
                out[fl] = page_key[4:]
            else:
                out[fl] = _ai1_field_hint(fl, section_id=sid, page_key=page_key)
        return out

    if sid == "keyword":
        if page_key.startswith(("SEO:", "TAG:")):
            return page_key[4:]
        return "キーワード"
    if sid == "reference_url":
        return "参考URL（ヒアリングにあれば）"
    if sid in {"faq_items", "review_items"}:
        return "ヒアリングに実本文がある場合のみ。なければ空。"
    return "ヒアリングから短く"


def build_ai1_result(
    blueprint: dict[str, Any],
    rows: list[dict[str, str]] | None = None,
    *,
    hearing: dict[str, Any] | None = None,
    site_name: str = "",
) -> OrderedDict[str, Any]:
    """AI-1 Result: {site, pages} with short nested hints (no id/label/mode/rule)."""
    bp = blueprint or {}
    pages = _page_index_from_blueprint(bp)
    pages_out: OrderedDict[str, OrderedDict[str, object]] = OrderedDict()

    if pages:
        for pk, page in pages.items():
            secs: OrderedDict[str, object] = OrderedDict()
            for sec in page.get("sections") or []:
                if not isinstance(sec, dict):
                    continue
                sid = str(sec.get("id") or "").strip()
                if not sid:
                    continue
                secs[sid] = _ai1_section_hint_value(sec, page_key=pk)
            pages_out[pk] = secs
    else:
        for row in rows or []:
            pk = page_export_key(
                group=str(row.get("group") or "page"),
                slug=str(row.get("page_slug") or ""),
                page_id=str(row.get("page_id") or ""),
                nav_label=str(row.get("nav_label") or ""),
            )
            sid = str(row.get("section_id") or "").strip()
            if not sid:
                continue
            pages_out.setdefault(pk, OrderedDict())
            if sid not in pages_out[pk]:
                pages_out[pk][sid] = _ai1_section_hint_value(
                    {"id": sid, "label": row.get("section_label") or ""},
                    page_key=pk,
                )

    return OrderedDict(
        [
            (
                "site",
                build_site_export_block(hearing, blueprint=bp, site_name=site_name),
            ),
            ("pages", pages_out),
        ]
    )


def _empty_section_value(sec: dict[str, Any], *, page_key: str = "") -> object:
    """Empty nested shell for AI-2 base (no planner hints)."""
    from ai_agent.v2.page_catalog import empty_nested_value, nested_fields_for_section

    sid = str(sec.get("id") or "").strip()
    fields = nested_fields_for_section(sec)
    if fields:
        shell = empty_nested_value(fields)
        if isinstance(shell, dict) and "keyword" in shell:
            if page_key.startswith("SEO:"):
                shell["keyword"] = page_key[4:]
            elif page_key.startswith("TAG:"):
                shell["keyword"] = page_key[4:]
        return shell
    if sid == "keyword":
        if page_key.startswith(("SEO:", "TAG:")):
            return page_key[4:]
        return ""
    return ""


def build_ai2_result(
    rows: list[dict[str, str]],
    blueprint: dict[str, Any] | None = None,
    *,
    hearing: dict[str, Any] | None = None,
    site_name: str = "",
) -> OrderedDict[str, Any]:
    """AI-2 Result: {site, pages} with filled nested content (same keys as AI-1)."""
    bp = blueprint or {}
    page_index = _page_index_from_blueprint(bp)
    page_order = list(page_index.keys())

    # Empty shells (not AI-1 hints) so missing copy stays blank
    pages_shell: OrderedDict[str, OrderedDict[str, object]] = OrderedDict()
    if page_index:
        for pk, page in page_index.items():
            secs: OrderedDict[str, object] = OrderedDict()
            for sec in page.get("sections") or []:
                if not isinstance(sec, dict):
                    continue
                sid = str(sec.get("id") or "").strip()
                if not sid:
                    continue
                secs[sid] = _empty_section_value(sec, page_key=pk)
            pages_shell[pk] = secs

    by_fill: dict[str, OrderedDict[str, object]] = {}
    order_seen: list[str] = []
    for row in rows or []:
        pk = page_export_key(
            group=str(row.get("group") or "page"),
            slug=str(row.get("page_slug") or ""),
            page_id=str(row.get("page_id") or ""),
            nav_label=str(row.get("nav_label") or ""),
        )
        sid = str(row.get("section_id") or "").strip()
        if not sid:
            continue
        if pk not in by_fill:
            by_fill[pk] = OrderedDict()
            order_seen.append(pk)
        body = _parse_body(row.get("text"))
        if body == "" or body is None:
            continue
        by_fill[pk][sid] = body

    pages_out: OrderedDict[str, OrderedDict[str, object]] = OrderedDict()
    for pk in page_order or list(pages_shell.keys()) or order_seen:
        base = OrderedDict(pages_shell.get(pk) or {})
        for sid, val in (by_fill.get(pk) or {}).items():
            base[sid] = val
        # Ensure section keys from fills even if not in shell
        if not base and pk in by_fill:
            base = OrderedDict(by_fill[pk])
        pages_out[pk] = base
    for pk in order_seen:
        if pk not in pages_out:
            pages_out[pk] = OrderedDict(by_fill.get(pk) or {})

    return OrderedDict(
        [
            (
                "site",
                build_site_export_block(hearing, blueprint=bp, site_name=site_name),
            ),
            ("pages", pages_out),
        ]
    )


_INDENT = "\u00A0" * 4  # NBSP × 4 — regular spaces collapse visually in Sheets/Excel wrap


def merge_outline_runs(runs: list[tuple[str, bool]]) -> list[tuple[str, bool]]:
    """Collapse consecutive same-bold runs (keeps Sheets textFormatRuns under limits)."""
    merged: list[tuple[str, bool]] = []
    for text, bold in runs or []:
        if not text:
            continue
        if merged and merged[-1][1] is bold:
            merged[-1] = (merged[-1][0] + text, bold)
        else:
            merged.append((text, bold))
    return merged


def format_result_outline(data: dict[str, Any] | None) -> str:
    """Human-readable Result outline: page/section keys first, 4-space nesting."""
    return "".join(text for text, _bold in iter_result_outline_runs(data or {}))


def iter_result_outline_runs(data: dict[str, Any] | None) -> list[tuple[str, bool]]:
    """Return merged (text, bold) runs for a Result outline (4× NBSP indent).

    Bold: site / pages / page names / section block keys.
    Leaf fields stay plain, indented under their section.
    """
    root = data if isinstance(data, dict) else {}
    runs: list[tuple[str, bool]] = []

    def add(text: str, *, bold: bool = False) -> None:
        if text:
            runs.append((text, bold))

    def pad(level: int) -> None:
        if level > 0:
            add(_INDENT * level)

    def line(level: int, label: str, *, bold: bool, value: str | None = None) -> None:
        pad(level)
        add(label, bold=bold)
        if value is not None:
            add(f": {value}")
        add("\n")

    def emit_scalar(level: int, key: str, value: Any, *, bold_key: bool) -> None:
        if isinstance(value, list):
            line(level, key, bold=bold_key)
            for item in value:
                pad(level + 1)
                add(f"- {item}\n")
            return
        if isinstance(value, dict):
            line(level, key, bold=bold_key)
            for sk, sv in value.items():
                emit_scalar(level + 1, str(sk), sv, bold_key=False)
            return
        line(level, key, bold=bold_key, value="" if value is None else str(value))

    site = root.get("site") if isinstance(root.get("site"), dict) else {}
    pages = root.get("pages") if isinstance(root.get("pages"), dict) else {}

    line(0, "site", bold=True)
    for key, val in site.items():
        # site fields are main items
        emit_scalar(1, str(key), val, bold_key=True)

    add("\n")
    line(0, "pages", bold=True)
    for page_name, sections in pages.items():
        line(1, str(page_name), bold=True)
        if not isinstance(sections, dict):
            pad(2)
            add(f"{sections}\n")
            continue
        for sec_key, sec_val in sections.items():
            emit_scalar(2, str(sec_key), sec_val, bold_key=True)
        add("\n")

    return merge_outline_runs(runs)


def outline_plain(data: dict[str, Any] | None) -> str:
    """Outline text for Sheet cell stringValue (NBSP indents preserved)."""
    return format_result_outline(data)


def build_pack_field_rows(
    *,
    test_name: str,
    site: str,
    model: str,
    model_note: str,
    hearing_file: str,
    hearing_url: str,
    system_prompt: str,
    result_json: str,
    result_outline: str = "",
) -> list[list[str]]:
    """Field / Value / Note body rows (no title / header)."""
    result_value = str(result_outline or result_json or "")
    return [
        ["Test name", test_name, ""],
        ["Site / client", site, ""],
        ["AI model", model, model_note],
        ["Input data (hearing file name)", hearing_file, ""],
        ["Input data (hearing URL)", hearing_url, ""],
        ["System prompt (FULL)", system_prompt, ""],
        ["Result", result_value, "Double-click cell · bold = page/section · indent = 4 spaces"],
    ]


def build_test_packs_payload(
    *,
    blueprint: dict[str, Any],
    rows: list[dict[str, str]],
    planner_prompt: str,
    writer_prompt: str,
    ai1_model: str = "",
    ai2_model: str = "",
    hearing_file: str = "",
    hearing_url: str = "",
    site_name: str = "",
    hearing: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build AI-1 / AI-2 pack sheets for xlsx + Google Sheets."""
    bp = blueprint or {}
    site = str(site_name or bp.get("site_name") or (rows[0].get("site") if rows else "") or "")
    type_label = _type_label(bp)
    stamp = datetime.now(UTC).strftime("%Y-%m-%d")
    ai1_struct = build_ai1_result(bp, rows, hearing=hearing, site_name=site)
    ai2_struct = build_ai2_result(rows, bp, hearing=hearing, site_name=site)
    ai1_json = json.dumps(ai1_struct, ensure_ascii=False, indent=4)
    ai2_json = json.dumps(ai2_struct, ensure_ascii=False, indent=4)
    ai1_outline = outline_plain(ai1_struct)
    ai2_outline = outline_plain(ai2_struct)

    ai1_title = f"Live — {type_label} — AI-1 Sections"
    ai2_title = f"Live — {type_label} — AI-2 Contents"
    ai1_test = f"Live export — {type_label} — AI-1 Sections — {stamp}"
    ai2_test = f"Live export — {type_label} — AI-2 Contents — {stamp}"

    ai1_fields = build_pack_field_rows(
        test_name=ai1_test,
        site=site,
        model=str(ai1_model or ""),
        model_note="AI-1 = Sections Planner",
        hearing_file=str(hearing_file or ""),
        hearing_url=str(hearing_url or ""),
        system_prompt=str(planner_prompt or ""),
        result_json=ai1_json,
        result_outline=ai1_outline,
    )
    ai2_fields = build_pack_field_rows(
        test_name=ai2_test,
        site=site,
        model=str(ai2_model or ai1_model or ""),
        model_note="AI-2 = Content Writer",
        hearing_file=str(hearing_file or ""),
        hearing_url=str(hearing_url or ""),
        system_prompt=str(writer_prompt or ""),
        result_json=ai2_json,
        result_outline=ai2_outline,
    )

    def sheet_values(title: str, fields: list[list[str]]) -> list[list[str]]:
        return [[title], ["Field", "Value", "Note"], *fields]

    return {
        "site_name": site,
        "type_label": type_label,
        "ai1": {
            "title": "AI-1 Sections",
            "banner": ai1_title,
            "values": sheet_values(ai1_title, ai1_fields),
            "result_json": ai1_json,
            "result_outline": ai1_outline,
            "result_runs": [[t, b] for t, b in iter_result_outline_runs(ai1_struct)],
        },
        "ai2": {
            "title": "AI-2 Contents",
            "banner": ai2_title,
            "values": sheet_values(ai2_title, ai2_fields),
            "result_json": ai2_json,
            "result_outline": ai2_outline,
            "result_runs": [[t, b] for t, b in iter_result_outline_runs(ai2_struct)],
        },
    }


def _outline_to_rich_text(runs: list[tuple[str, bool]] | list[list[Any]]):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    bold_font = InlineFont(b=True, rFont="Consolas", sz=10)
    plain_font = InlineFont(b=False, rFont="Consolas", sz=10)
    blocks: list[TextBlock] = []
    for item in runs or []:
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            text, bold = str(item[0] or ""), bool(item[1])
        else:
            continue
        if not text:
            continue
        blocks.append(TextBlock(bold_font if bold else plain_font, text))
    if not blocks:
        return ""
    return CellRichText(*blocks)


def packs_to_xlsx_bytes(packs: dict[str, Any]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for spreadsheet export") from exc

    wb = Workbook()
    first = True
    for key in ("ai1", "ai2"):
        pack = packs.get(key) or {}
        values = pack.get("values") or []
        title = str(pack.get("title") or key)[:31]
        if first:
            ws = wb.active
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title)
        for row in values:
            ws.append(list(row))
        if values:
            ws["A1"].font = Font(bold=True, size=13)
            for cell in ws[2]:
                cell.font = Font(bold=True)
        # Result cell: bold page/section keys + NBSP indent (row 9 = Result)
        result_runs = pack.get("result_runs") or []
        result_cell = ws.cell(9, 2)
        if result_runs:
            result_cell.value = _outline_to_rich_text(result_runs)
        result_cell.font = Font(name="Consolas", size=10)
        result_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 80
        ws.column_dimensions["C"].width = 28
        # Wrap prompt + result
        for row_idx in (8, 9):
            cell = ws.cell(row_idx, 2)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row_idx].height = 280

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sections_to_xlsx_bytes(
    rows: list[dict[str, str]],
    *,
    blueprint: dict[str, Any] | None = None,
    planner_prompt: str = "",
    writer_prompt: str = "",
    ai1_model: str = "",
    ai2_model: str = "",
    hearing_file: str = "",
    hearing_url: str = "",
    site_name: str = "",
    hearing: dict[str, Any] | None = None,
) -> bytes:
    """Live xlsx = Test1-style AI-1 + AI-2 packs (Field / Value / Note)."""
    packs = build_test_packs_payload(
        blueprint=blueprint or {},
        rows=rows,
        planner_prompt=planner_prompt,
        writer_prompt=writer_prompt,
        ai1_model=ai1_model,
        ai2_model=ai2_model,
        hearing_file=hearing_file,
        hearing_url=hearing_url,
        site_name=site_name,
        hearing=hearing,
    )
    return packs_to_xlsx_bytes(packs)
